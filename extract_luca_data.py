import openpyxl
import json
import csv
import os
import re

def normalize(text):
    if text is None: return ""
    return re.sub(r'\s+', ' ', str(text).strip())

def extract_data():
    wb_path = r"c:\Users\james\Documents\Code\LUCA\Land Use Change Assessor - LUCA (V3).xlsm"
    print(f"Loading workbook: {wb_path}")
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    
    # --- Part 1: Assumptions Metadata ---
    ws_assumptions = wb['Assumptions']
    
    catchments_ordered = []
    # Read row 12 for catchments
    row_12 = [cell.value for cell in ws_assumptions[12]][4:] 
    for c in row_12:
        nc = normalize(c)
        if nc and nc not in catchments_ordered:
            catchments_ordered.append(nc)
            
    # Read row 13 for predefined land uses (take the first catchment's 22 columns)
    predefined_land_uses = []
    row_13 = [cell.value for cell in ws_assumptions[13]][4:26]
    for lu in row_13:
        if lu:
            predefined_land_uses.append(normalize(lu))
            
    print(f"Extracted {len(catchments_ordered)} catchments.")
    print(f"Extracted {len(predefined_land_uses)} predefined land uses.")
    
    # Indicators
    reported_indicator_names = [
        "Value-added - Direct impact",
        "Value-added - Direct & indirect impacts",
        "Value-added - Direct, indirect & induced impacts",
        "Employment - Direct impact",
        "Employment - Direct & indirect impacts",
        "Employment - Direct, indirect & induced impacts",
        "Land value",
        "Income",
        "N loss",
        "P loss",
        "Sediment loss",
        "E.coli loss",
        "River water quality",
        "Freshwater invertebrates",
        "Estuarine health",
        "Groundwater quality",
        "Pressure on water quantity",
        "Soil quality",
        "Biodiversity",
        "GHG emissions",
        "Connection to nature",
        "Access to basic amenities",
        "Housing affordability",
        "Life satisfaction (current)",
        "Life satisfaction (future)",
        "Sense of belonging"
    ]
    
    glossary = {}
    if 'Glossary' in wb.sheetnames:
        ws_glossary = wb['Glossary']
        for row in ws_glossary.iter_rows(min_row=3, max_row=35, values_only=True):
            ind = normalize(row[2])
            desc = normalize(row[3])
            if ind and desc:
                glossary[ind.lower()] = desc
                
    indicators = []
    # Map indicator names to Assumptions sheet rows (14 to 39)
    # We'll just assume they are in order as per reported_indicator_names
    # but let's be safer and match by name in column 3
    row_map = {}
    for r in range(14, 40):
        name = normalize(ws_assumptions.cell(row=r, column=3).value)
        if name: row_map[name] = r

    for name in reported_indicator_names:
        r = row_map.get(normalize(name))
        if not r: continue
        
        domain = normalize(ws_assumptions.cell(row=r, column=2).value)
        # Handle merged domain cells
        if not domain:
            for prev_r in range(r, 13, -1):
                d = normalize(ws_assumptions.cell(row=prev_r, column=2).value)
                if d:
                    domain = d
                    break
        
        unit = normalize(ws_assumptions.cell(row=r, column=4).value)
        if name == "Income": unit = "$/ha"
        
        is_scoring = "scoring" in unit.lower() or domain == "Social" or normalize(name) in [
            "Sediment loss", "E.coli loss", "River water quality", "Freshwater invertebrates",
            "Estuarine health", "Groundwater quality", "Pressure on water quantity",
            "Soil quality", "Biodiversity"
        ]
        
        indicators.append({
            'domain': domain,
            'name': name,
            'unit': unit,
            'type': "scoring" if is_scoring else "metric",
            'description': glossary.get(name.lower(), "No description available.")
        })

    # --- Part 2: Baseline Extraction ---
    ws_baseline = wb['LU-Baseline']
    baseline_data = {c: {} for c in catchments_ordered}
    
    # Catchment headers are at Row 5
    baseline_headers = [normalize(cell.value) for cell in ws_baseline[5]]
    catchment_cols = {}
    for idx, h in enumerate(baseline_headers):
        if h in catchments_ordered:
            catchment_cols[h] = idx + 1
            
    print(f"Matched {len(catchment_cols)} catchments in LU-Baseline.")

    # Data is at Rows 6 to 30
    for r in range(6, 31):
        lu_raw = ws_baseline.cell(row=r, column=1).value
        if not lu_raw: continue
        lu_name = normalize(lu_raw)
        for c, col_idx in catchment_cols.items():
            val = ws_baseline.cell(row=r, column=col_idx).value
            try:
                val_float = float(val) if val is not None and str(val).strip() != "" else 0.0
            except ValueError:
                val_float = 0.0
            baseline_data[c][lu_name] = val_float

    custom_land_uses = ["Custom Land Use A", "Custom Land Use B", "Custom Land Use C"]
    all_land_uses = predefined_land_uses + custom_land_uses

    # --- Part 3: Assumptions Extraction ---
    assumptions_data = {c: {} for c in catchments_ordered}
    
    # Map (Catchment, LandUse) to column index in Assumptions (Row 12, 13)
    col_map = {}
    for col_idx in range(5, ws_assumptions.max_column + 1):
        c_name = normalize(ws_assumptions.cell(row=12, column=col_idx).value)
        lu_name = normalize(ws_assumptions.cell(row=13, column=col_idx).value)
        if c_name and lu_name:
            col_map[(c_name, lu_name)] = col_idx

    row_emp_direct = row_map.get(normalize("Employment - Direct impact"), 17)
    row_income_raw = row_map.get(normalize("Income"), 21)

    for c in catchments_ordered:
        for ind in indicators:
            ind_name = ind['name']
            if ind_name not in assumptions_data[c]: assumptions_data[c][ind_name] = {}
            r = row_map.get(normalize(ind_name))
            
            for lu in predefined_land_uses:
                col_idx = col_map.get((c, lu))
                val = 0.0
                if col_idx and r:
                    if ind_name == "Income":
                        inc_val = ws_assumptions.cell(row=row_income_raw, column=col_idx).value
                        emp_val = ws_assumptions.cell(row=row_emp_direct, column=col_idx).value
                        try:
                            inc_f = float(inc_val) if inc_val is not None and str(inc_val).strip() != "" else 0.0
                            emp_f = float(emp_val) if emp_val is not None and str(emp_val).strip() != "" else 0.0
                            val = inc_f * emp_f
                        except: val = 0.0
                    else:
                        raw_val = ws_assumptions.cell(row=r, column=col_idx).value
                        try:
                            val = float(raw_val) if raw_val is not None and str(raw_val).strip() != "" else 0.0
                        except: val = 0.0
                assumptions_data[c][ind_name][lu] = val
            for clu in custom_land_uses:
                assumptions_data[c][ind_name][clu] = 0.0

    # --- Part 4: Save ---
    master_object = {
        'catchments': catchments_ordered,
        'predefinedLandUses': predefined_land_uses,
        'customLandUses': custom_land_uses,
        'allLandUses': all_land_uses,
        'indicators': indicators,
        'baseline': baseline_data,
        'assumptions': assumptions_data
    }
    
    base_dir = r"c:\Users\james\Documents\Code\LUCA"
    with open(os.path.join(base_dir, "data.js"), "w", encoding="utf-8") as f:
        f.write("const LUCA_DATA = " + json.dumps(master_object, indent=2, ensure_ascii=False) + ";")
    
    with open(os.path.join(base_dir, "baseline.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Land Use"] + catchments_ordered)
        for lu in all_land_uses:
            writer.writerow([lu] + [baseline_data[c].get(lu, 0.0) for c in catchments_ordered])
            
    print(f"Sample Check (Waihī Estuary - Sheep + Beef SEPP): {baseline_data.get('Waihī Estuary', {}).get('Sheep + Beef (SEPP)', 'MISSING')}")
    print("Done!")

if __name__ == "__main__":
    extract_data()
